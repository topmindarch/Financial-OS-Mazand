from io import BytesIO
from typing import Any
from openpyxl import Workbook, load_workbook

HEADERS = ["product", "side", "group", "mode", "price", "factor", "fixed_delta"]
PERSIAN_HEADERS = {"محصول":"product", "جهت":"side", "گروه":"group", "روش":"mode", "قیمت":"price", "ضریب":"factor", "تعدیل ثابت":"fixed_delta"}

def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()

def _side(value: str) -> str:
    v = value.strip().lower()
    if v in {"buy", "خرید", "خري", "خریدن"}: return "buy"
    if v in {"sell", "فروش", "فرو"}: return "sell"
    return ""

def parse_workbook(payload: bytes) -> tuple[list[dict], list[str]]:
    errors: list[str] = []
    rows: list[dict] = []
    try:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        raw_headers = next(iterator, None)
        if not raw_headers: return [], ["ردیف عنوان ستون‌ها وجود ندارد."]
        headers = [PERSIAN_HEADERS.get(_text(raw).lower(), _text(raw).lower()) for raw in raw_headers]
        missing = [h for h in HEADERS if h not in headers]
        if missing: return [], [f"ستون‌های لازم پیدا نشدند: {', '.join(missing)}"]
        positions = {h: headers.index(h) for h in HEADERS}
        for excel_row, values in enumerate(iterator, start=2):
            if not any(v is not None and str(v).strip() for v in values): continue
            def cell(name: str) -> str:
                index = positions[name]
                return _text(values[index] if index < len(values) else None)
            product, side, group, mode = cell("product"), _side(cell("side")), cell("group"), cell("mode").lower()
            price_text = cell("price").replace(",", "").replace("٬", "")
            factor_text = cell("factor") or "1"
            delta_text = cell("fixed_delta").replace(",", "") or "0"
            try:
                price = int(float(price_text)) if price_text else None
                factor = float(factor_text)
                fixed_delta = int(float(delta_text))
            except ValueError:
                errors.append(f"ردیف {excel_row}: قیمت، ضریب یا تعدیل عددی نیست."); continue
            if not product: errors.append(f"ردیف {excel_row}: نام محصول خالی است."); continue
            if not side: errors.append(f"ردیف {excel_row}: جهت باید buy/sell یا خرید/فروش باشد."); continue
            if price is not None and price < 0: errors.append(f"ردیف {excel_row}: قیمت منفی مجاز نیست."); continue
            if group:
                if mode not in {"direct", "multiplier"}: errors.append(f"ردیف {excel_row}: روش گروه باید direct یا multiplier باشد."); continue
                if mode == "direct" and price is None: errors.append(f"ردیف {excel_row}: برای گروه direct قیمت لازم است."); continue
                if mode == "multiplier" and factor <= 0: errors.append(f"ردیف {excel_row}: ضریب باید بزرگ‌تر از صفر باشد."); continue
            elif price is None:
                errors.append(f"ردیف {excel_row}: برای قیمت پایه، price باید پر و group خالی باشد."); continue
            rows.append({"excel_row":excel_row,"product":product,"side":side,"group":group or None,"mode":mode if group else None,"price":price,"factor":factor,"fixed_delta":fixed_delta})
        workbook.close()
    except Exception as exc:
        return [], [f"فایل Excel قابل خواندن نیست: {exc}"]
    return rows, errors

def make_template() -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Prices"
    sheet.append(HEADERS)
    sheet.append(["پرتقال نمونه", "buy", "", "", 250000, 1, 0])
    sheet.append(["پرتقال نمونه", "buy", "گروه نمونه", "direct", 260000, 1, 0])
    sheet.append(["نارنگی نمونه", "sell", "گروه ضریبی", "multiplier", "", 1.08, -5000])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:G4"
    for column, width in {"A":28,"B":12,"C":24,"D":16,"E":18,"F":14,"G":16}.items(): sheet.column_dimensions[column].width = width
    output = BytesIO(); workbook.save(output); output.seek(0)
    return output
